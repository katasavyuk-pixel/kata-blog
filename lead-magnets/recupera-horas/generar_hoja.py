"""
La hoja de tu semana: half of the signup gift «Recupera horas de tu semana con IA»
(Kata, 2026-10-16: «algo que sirva a todos, tengan negocio o no»).

The reader lists the repetitive tasks of their week, work or personal; the sheet
gives back the hours a year each one takes, how many an AI could give back and
where to start. With the value of an hour (theirs, or whoever does the task) it
also gives euros, which is where a business owner comes in.

Every result is a formula, so the sheet recalculates as it is filled in. The
only fixed numbers are the settings on «Empieza aquí» (weeks a year, hours a
week, what «casi toda / la mitad / poco» mean), each one written next to the
reason for it and editable.

    python3 lead-magnets/recupera-horas/generar_hoja.py
    python3 ~/.claude/skills/xlsx/scripts/recalc.py lead-magnets/recupera-horas/hoja-de-tu-semana.xlsx
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SALIDA = Path(__file__).with_name("hoja-de-tu-semana.xlsx")

# The site's palette (src/app/globals.css): ink, paper, red.
TINTA = "140A0D"
PAPEL = "F2EDEB"
ROJO = "E11423"
ROJO_TEXTO = "C40F1F"
GRIS = "5E4B50"
RELLENAR = "FFF2CC"  # the cells the reader fills in

FUENTE = "Arial"
f = lambda **k: Font(name=FUENTE, **k)  # noqa: E731

INICIO = "Empieza aquí"
SEMANA = "Tu semana"
EJEMPLO = "Ejemplo"
REF = f"'{INICIO}'"

# Settings on «Empieza aquí»: row numbers fixed below, referenced by the formulas.
FILA_SEMANAS = 21
FILA_HORAS_SEMANA = 22
FILA_VALOR_HORA = 23
FILAS_NIVELES = (27, 30)  # «Casi toda» … «Nada»
NIVELES = [("Casi toda", 0.8), ("La mitad", 0.5), ("Poco", 0.2), ("Nada", 0)]

PRIMERA = 12  # first task row on the week sheets
ULTIMA = 41  # 30 rows
CABECERA = PRIMERA - 1

borde = Side(style="thin", color="E3D9D6")
celda = Border(left=borde, right=borde, top=borde, bottom=borde)
regla = Border(top=Side(style="medium", color=ROJO))


def ajustes(ws):
    """«Empieza aquí»: what the sheet is, how to fill it, the settings and what to do next."""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 62

    ws["B2"] = "Recupera horas de tu semana con IA"
    ws["B2"].font = f(size=20, bold=True, color=TINTA)
    ws["B3"] = "La hoja de tu semana · Kata Ivanovych · kata.ianexora.com"
    ws["B3"].font = f(size=10, color=GRIS)

    ws["B5"] = "Cómo se usa (15 minutos)"
    ws["B5"].font = f(size=12, bold=True, color=ROJO_TEXTO)
    pasos = [
        "1. Ve a la pestaña «Tu semana» y apunta tus tareas repetitivas: las que haces casi igual cada vez. Correos parecidos, pasar datos de un sitio a otro, resúmenes, informes, buscar algo que ya tenías. Del trabajo o de tu vida.",
        "2. En cada una, pon cuántos minutos te lleva y cuántas veces la haces a la semana.",
        "3. Elige cuánto de esa tarea podría hacer una IA: casi toda, la mitad, poco o nada. Es tu estimación; si dudas, elige menos.",
        "4. La hoja te dice las horas al año que se te van en cada tarea, las que te puede devolver la IA y por cuál empezar.",
        "5. Opcional: pon abajo cuánto vale una hora tuya (o en la tabla, la de quien hace cada tarea) y lo verás también en euros.",
    ]
    for i, paso in enumerate(pasos):
        c = ws.cell(row=6 + i, column=2, value=paso)
        ws.merge_cells(start_row=6 + i, start_column=2, end_row=6 + i, end_column=4)
        c.font = f(size=11, color=TINTA)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[6 + i].height = 32

    ws["B12"] = "Las celdas amarillas son las que rellenas tú. El resto se calcula solo."
    ws["B12"].fill = PatternFill("solid", fgColor=RELLENAR)
    ws["B12"].font = f(size=10, color=TINTA)
    ws["B13"] = "Mira la pestaña «Ejemplo» para ver una semana ya rellena. La hoja es tuya: no envía nada a ningún sitio."
    ws["B13"].font = f(size=10, color=GRIS)
    ws.merge_cells("B13:D13")

    ws["B19"] = "Ajustes"
    ws["B19"].font = f(size=12, bold=True, color=ROJO_TEXTO)
    ws["B20"] = "Puedes cambiarlos. Son los números que la hoja usa para calcular."
    ws["B20"].font = f(size=10, color=GRIS)
    filas = [
        (FILA_SEMANAS, "Semanas que trabajas al año", 45, "Un año tiene 52. Unas 4 se van en vacaciones y unas 3 en festivos. Pon las tuyas.", "0"),
        (FILA_HORAS_SEMANA, "Horas de tu semana de trabajo", 40, "Para contarte las horas recuperadas en semanas de trabajo.", "0"),
        (FILA_VALOR_HORA, "Cuánto vale una hora tuya (€, opcional)", None, "Lo que cobras o lo que te cuesta una hora. Si lo dejas vacío, la hoja solo cuenta horas.", '#,##0 "€"'),
    ]
    for fila, nombre, valor, nota, formato in filas:
        ws.cell(row=fila, column=2, value=nombre).font = f(size=11, color=TINTA)
        c = ws.cell(row=fila, column=3, value=valor)
        c.fill = PatternFill("solid", fgColor=RELLENAR)
        c.font = f(size=11, bold=True, color=TINTA)
        c.number_format = formato
        c.border = celda
        nota_c = ws.cell(row=fila, column=4, value=nota)
        nota_c.font = f(size=9, color=GRIS)
        nota_c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[fila].height = 26

    ws["B25"] = "Cuánto puede hacer la IA"
    ws["B25"].font = f(size=11, bold=True, color=TINTA)
    ws["D25"] = "Son estimaciones para ordenar tus tareas, no una promesa. Si crees que «la mitad» es menos, cámbialo."
    ws["D25"].font = f(size=9, color=GRIS)
    ws["D25"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[25].height = 26
    ws["B26"] = "Si eliges…"
    ws["C26"] = "…te devuelve"
    for celda_t in (ws["B26"], ws["C26"]):
        celda_t.font = f(size=9, bold=True, color=GRIS)
    for i, (nombre, parte) in enumerate(NIVELES):
        fila = FILAS_NIVELES[0] + i
        ws.cell(row=fila, column=2, value=nombre).font = f(size=11, color=TINTA)
        c = ws.cell(row=fila, column=3, value=parte)
        c.number_format = "0%"
        c.fill = PatternFill("solid", fgColor=RELLENAR)
        c.font = f(size=11, color=TINTA)
        c.border = celda

    ws["B33"] = "Y después"
    ws["B33"].font = f(size=12, bold=True, color=ROJO_TEXTO)
    despues = [
        "Empieza por la tarea número 1 de «Tu semana». Abre tu IA (ChatGPT, Claude o Gemini) con tu ficha puesta y pásale este mensaje, cambiando lo que va entre corchetes:",
        "«Hago esta tarea [cada cuánto]: [la tarea]. Me lleva [minutos]. Lo que me llega es [qué recibes: un correo, una factura, unas notas] y lo que tiene que salir es [qué entregas]. ¿Qué parte puedes hacer tú, y cómo te lo paso cada vez para que lo hagas igual?»",
        "Haz la tarea con ella esta semana. Si te ahorra tiempo, pasa a la número 2. Si no, cuéntame qué falló respondiendo a cualquiera de mis correos: lo leo yo.",
    ]
    for i, texto in enumerate(despues):
        fila = 34 + i
        c = ws.cell(row=fila, column=2, value=texto)
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
        c.font = f(size=11, color=TINTA, italic=(i == 1))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.fill = PatternFill("solid", fgColor=PAPEL)
        ws.row_dimensions[fila].height = 46 if i == 1 else 32


def semana(ws, tareas=(), titulo="Tu semana", nota=None):
    """The table. `tareas` pre-fills rows (the example); the reader's sheet starts empty."""
    ws.sheet_view.showGridLines = False
    anchos = {"A": 56, "B": 13, "C": 13, "D": 22, "E": 13, "F": 15, "G": 13, "H": 16, "I": 16, "J": 4}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho

    ws["A1"] = titulo
    ws["A1"].font = f(size=18, bold=True, color=TINTA)
    if nota:
        ws["A2"] = nota
        ws["A2"].font = f(size=10, italic=True, color=ROJO_TEXTO)

    rango = lambda col: f"${col}${PRIMERA}:${col}${ULTIMA}"  # noqa: E731
    resumen = [
        ("Horas al año en tareas repetitivas", f"=SUM({rango('E')})", "#,##0"),
        ("Horas que te puede devolver la IA", f"=SUM({rango('F')})", "#,##0"),
        ("Eso son, en semanas de trabajo", f"=IF(B5=0,0,B5/{REF}!$C${FILA_HORAS_SEMANA})", "#,##0.0"),
        ("Y en euros (si has puesto el valor de la hora)", f'=IF(COUNT({rango("I")})=0,"—",SUM({rango("I")}))', '#,##0 "€"'),
    ]
    for i, (nombre, formula, formato) in enumerate(resumen):
        fila = 4 + i
        ws.cell(row=fila, column=1, value=nombre).font = f(size=11, color=TINTA)
        c = ws.cell(row=fila, column=2, value=formula)
        c.font = f(size=12, bold=True, color=TINTA)
        c.number_format = formato
        c.alignment = Alignment(horizontal="right")

    ws["D4"] = "Por dónde empezar"
    ws["D4"].font = f(size=11, bold=True, color=ROJO_TEXTO)
    for i in range(3):
        fila = 5 + i
        tarea = f'=IFERROR(INDEX($A${PRIMERA}:$A${ULTIMA},MATCH(LARGE($J${PRIMERA}:$J${ULTIMA},{i + 1}),$J${PRIMERA}:$J${ULTIMA},0)),"")'
        horas = f'=IFERROR(INDEX($F${PRIMERA}:$F${ULTIMA},MATCH(LARGE($J${PRIMERA}:$J${ULTIMA},{i + 1}),$J${PRIMERA}:$J${ULTIMA},0)),"")'
        ws.cell(row=fila, column=4, value=f"{i + 1}.").font = f(size=11, bold=True, color=TINTA)
        c = ws.cell(row=fila, column=5, value=tarea)
        c.font = f(size=11, color=TINTA)
        ws.merge_cells(start_row=fila, start_column=5, end_row=fila, end_column=8)
        h = ws.cell(row=fila, column=9, value=horas)
        h.number_format = '#,##0 "h/año"'
        h.font = f(size=10, color=GRIS)

    cabeceras = [
        "Tarea",
        "Minutos cada vez",
        "Veces por semana",
        "¿Cuánto puede hacer la IA?",
        "Horas al año",
        "Horas que te devuelve al año",
        "Orden",
        "€ por hora de quien la hace (opcional)",
        "€ al año que te devuelve",
    ]
    for col, texto in enumerate(cabeceras, start=1):
        c = ws.cell(row=CABECERA, column=col, value=texto)
        c.font = f(size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=TINTA)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if col > 1 else "left")
    ws.row_dimensions[CABECERA].height = 42
    ws.freeze_panes = ws.cell(row=PRIMERA, column=1)

    lista = DataValidation(type="list", formula1='"' + ",".join(n for n, _ in NIVELES) + '"', allow_blank=True)
    lista.error = "Elige una de la lista: Casi toda, La mitad, Poco o Nada."
    lista.errorTitle = "¿Cuánto puede hacer la IA?"
    ws.add_data_validation(lista)

    niveles = f"{REF}!$B${FILAS_NIVELES[0]}:$B${FILAS_NIVELES[1]}"
    partes = f"{REF}!$C${FILAS_NIVELES[0]}:$C${FILAS_NIVELES[1]}"
    for fila in range(PRIMERA, ULTIMA + 1):
        dato = tareas[fila - PRIMERA] if fila - PRIMERA < len(tareas) else (None, None, None, None, None)
        nombre, minutos, veces, nivel, euros = dato
        for col, valor in ((1, nombre), (2, minutos), (3, veces), (4, nivel), (8, euros)):
            c = ws.cell(row=fila, column=col, value=valor)
            c.fill = PatternFill("solid", fgColor=RELLENAR)
            c.border = celda
            c.font = f(size=11, color=TINTA)
        lista.add(f"D{fila}")
        ws.cell(row=fila, column=8).number_format = '#,##0 "€"'

        formulas = {
            # Hours a year: minutes × times a week × working weeks.
            5: (f'=IF(OR(B{fila}="",C{fila}=""),"",B{fila}*C{fila}*{REF}!$C${FILA_SEMANAS}/60)', "#,##0"),
            # What the AI gives back: those hours × the share chosen for the level.
            6: (f'=IF(E{fila}="","",E{fila}*IFERROR(INDEX({partes},MATCH(D{fila},{niveles},0)),0))', "#,##0"),
            7: (f'=IF(OR(F{fila}="",F{fila}=0),"",RANK(F{fila},$F${PRIMERA}:$F${ULTIMA}))', "0"),
            # Euros: the row's own rate, or the one on «Empieza aquí».
            9: (
                f'=IF(OR(F{fila}="",AND(H{fila}="",{REF}!$C${FILA_VALOR_HORA}="")),"",F{fila}*IF(H{fila}<>"",H{fila},{REF}!$C${FILA_VALOR_HORA}))',
                '#,##0 "€"',
            ),
            # Hidden: a unique key so «Por dónde empezar» never repeats a task on a tie.
            10: (f'=IF(OR(F{fila}="",F{fila}=0),"",F{fila}+ROW()/100000)', "General"),
        }
        for col, (formula, formato) in formulas.items():
            c = ws.cell(row=fila, column=col, value=formula)
            c.number_format = formato
            c.border = celda
            c.font = f(size=11, color=TINTA, bold=(col == 6))
    ws.column_dimensions["J"].hidden = True

    total = ULTIMA + 1
    ws.cell(row=total, column=1, value="Total").font = f(size=11, bold=True, color=TINTA)
    for col, formato in ((5, "#,##0"), (6, "#,##0"), (9, '#,##0 "€"')):
        letra = "EFI"[(5, 6, 9).index(col)]
        c = ws.cell(row=total, column=col, value=f"=SUM({letra}{PRIMERA}:{letra}{ULTIMA})")
        c.number_format = formato
        c.font = f(size=11, bold=True, color=TINTA)
    for col in range(1, 10):
        ws.cell(row=total, column=col).border = regla


def imprimible(ws, horizontal):
    """One page wide when printed: nobody prints a table cut down the middle."""
    ws.page_setup.orientation = "landscape" if horizontal else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    libro = Workbook()
    inicio = libro.active
    inicio.title = INICIO
    ajustes(inicio)
    semana(libro.create_sheet(SEMANA))
    # Invented on purpose, to show the format: a week with work and home in it.
    semana(
        libro.create_sheet(EJEMPLO),
        tareas=[
            ("Contestar correos que se repiten (presupuestos, dudas de clientes)", 10, 15, "La mitad", 20),
            ("Pasar datos de facturas o pedidos a una hoja", 5, 20, "Casi toda", 20),
            ("Preparar el informe de la semana", 60, 1, "La mitad", 20),
            ("Resumir reuniones y repartir las tareas", 20, 3, "Casi toda", 20),
            ("Buscar un documento o un dato que ya tenía", 5, 10, "Poco", 20),
            ("Planificar las comidas y la compra de la semana", 30, 1, "La mitad", None),
        ],
        titulo="Ejemplo",
        nota="Una semana inventada, para que veas cómo se rellena. La tuya va en la pestaña «Tu semana».",
    )
    imprimible(inicio, horizontal=False)
    imprimible(libro[SEMANA], horizontal=True)
    imprimible(libro[EJEMPLO], horizontal=True)
    libro.active = 0
    libro.save(SALIDA)
    print(SALIDA)


if __name__ == "__main__":
    main()
